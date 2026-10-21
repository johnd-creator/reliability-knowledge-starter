"""Passive Excel reader: raw scalars, source serials and stored formula cache only."""
from pathlib import Path
import hashlib, datetime, zipfile, struct, xml.etree.ElementTree as ET
import openpyxl, xlrd, olefile
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
digest=lambda b:hashlib.sha256(b).hexdigest()
N={"s":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
class Reader:
 def __init__(self,p):
  p=Path(p)
  if p.is_symlink() or p.suffix.lower() not in {'.xls','.xlsx'} or not 0<p.stat().st_size<=32*1024*1024:raise ValueError('UNSUPPORTED_OR_OVERSIZED_WORKBOOK')
  if p.suffix.lower()=='.xlsx':
   with zipfile.ZipFile(p) as archive:
    members=archive.infolist()
    if len(members)>10000 or sum(x.file_size for x in members)>256*1024*1024 or any(x.file_size>32*1024*1024 for x in members):raise ValueError('WORKBOOK_EXPANSION_LIMIT')
    if any('vbaProject' in x.filename for x in members):raise ValueError('MACRO_WORKBOOK_NOT_SUPPORTED')
  self.path=p;self.sha=digest(p.read_bytes());self.xls=p.suffix.lower()=='.xls';self.raw={};self.formulas={}
  if self.xls:
   self.w=xlrd.open_workbook(p,on_demand=True,formatting_info=True);self.names=self.w.sheet_names()
   with olefile.OleFileIO(p) as o:
    data=o.openstream('Workbook' if o.exists('Workbook') else 'Book').read()
   offsets=[];at=0
   while at+4<=len(data):
    typ,n=struct.unpack_from('<HH',data,at);payload=data[at+4:at+4+n]
    if typ==0x85 and len(payload)>=4:offsets.append(struct.unpack_from('<I',payload)[0])
    at+=4+n
   for i,start in enumerate(offsets):
    at=start;found={}
    while at+4<=len(data):
     typ,n=struct.unpack_from('<HH',data,at);payload=data[at+4:at+4+n]
     if typ==6 and n>=22:
      r,c=struct.unpack_from('<HH',payload);found[get_column_letter(c+1)+str(r+1)]=digest(payload[20:])
     if typ==0x0A:break
     at+=4+n
    self.formulas[i]=found
  else:
   self.w=openpyxl.load_workbook(p,read_only=False,data_only=True,keep_links=False);self.names=self.w.sheetnames
   with zipfile.ZipFile(p) as z:
    wb=ET.fromstring(z.read('xl/workbook.xml'));rels=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
    targets={r.attrib['Id']:r.attrib['Target'] for r in rels}
    for i,sh in enumerate(wb.find('s:sheets',N)):
     target=targets[sh.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
     target=target.lstrip('/') if target.startswith('/') else 'xl/'+target
     cells={};regions=[]
     for c in ET.fromstring(z.read(target)).findall('.//s:sheetData/s:row/s:c',N):
      v=c.find('s:v',N);f=c.find('s:f',N)
      cells[c.attrib['r']]={'xml_type':c.get('t'),'xml_value':None if v is None else v.text,'formula':None if f is None else f.text,'formula_present':f is not None}
      if f is not None and f.get('ref'):regions.append((f.get('ref'),f.text))
     for ref,formula in regions:
      a,b,c,d=range_boundaries(ref)
      if (c-a+1)*(d-b+1)>100000:raise ValueError('FORMULA_REGION_LIMIT')
      for row in range(b,d+1):
       for col in range(a,c+1):
        item=cells.setdefault(get_column_letter(col)+str(row),{})
        item['formula_present']=True
        if item.get('formula') is None:item['formula']=formula
     self.raw[i]=cells
 def shape(self,i):
  s=self.w.sheet_by_index(i) if self.xls else self.w.worksheets[i]
  return (s.nrows,s.ncols) if self.xls else (s.max_row,s.max_column)
 def cell(self,i,address):
  r,c=coordinate_to_tuple(address)
  if self.xls:
   sh=self.w.sheet_by_index(i)
   cell=sh.cell(r-1,c-1) if r<=sh.nrows and c<=sh.ncols else None
   raw=None if cell is None or cell.ctype in [0,6] else cell.value
   form=address in self.formulas.get(i,{})
   date=None
   if cell and cell.ctype==3:
    dt=xlrd.xldate.xldate_as_datetime(raw,self.w.datemode)
    date=dt.isoformat()
   fmt=self.w.format_map[self.w.xf_list[cell.xf_index].format_key].format_str if cell else None
   return {'cell':address,'raw_value':raw,'date_interpretation':date,'source_type':None if cell is None else cell.ctype,'formula_present':form,'formula_sha256':self.formulas.get(i,{}).get(address),'cache_status':'STORED_FORMULA_RESULT_NOT_RECALCULATED' if form else 'LITERAL','number_format':fmt,'xml_value':None}
  cell=self.w.worksheets[i][address];v=cell.value;meta=self.raw[i].get(address,{})
  date=v.isoformat() if isinstance(v,(datetime.date,datetime.datetime,datetime.time)) else None
  raw=v
  if date:
   raw=meta.get('xml_value')
   if meta.get('xml_type')=='n' or meta.get('xml_type') is None:
    try:raw=float(raw);raw=int(raw) if raw.is_integer() else raw
    except (TypeError,ValueError):pass
  return {'cell':address,'raw_value':raw,'date_interpretation':date,'source_type':cell.data_type,'formula_present':meta.get('formula_present',False),'formula_sha256':digest(meta['formula'].encode()) if meta.get('formula') else None,'cache_status':('STORED_FORMULA_RESULT_NOT_RECALCULATED' if v is not None else 'NO_STORED_FORMULA_RESULT') if meta.get('formula_present') else 'LITERAL','number_format':cell.number_format,'xml_value':meta.get('xml_value')}
 def close(self):self.w.release_resources() if self.xls else self.w.close()
