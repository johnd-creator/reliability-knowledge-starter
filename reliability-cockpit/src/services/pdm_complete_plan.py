"""Passive complete source plan. Source occurrences are not canonical measurements.

Every populated physical cell (including a formula without a cached result) has
one terminal coverage category. Unknown regions remain navigable RAW rows.
Original files never change; classification versions are immutable evidence.
"""
from collections import Counter, defaultdict
from datetime import date
from hashlib import sha256
from pathlib import Path
import base64, json, math, os, struct, zipfile, xml.etree.ElementTree as ET
import olefile
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from src.adapters.pdm_workbooks import Reader
from src.services.pdm_history_staging import load_catalog, iter_events, event_key

PARSER = 'NADI_COMPLETE_SOURCE_V1'
STREAMS = ('sources', 'profiles', 'cells', 'events')


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False)


def digest(value):
    return sha256(canonical(value).encode()).hexdigest()


def scalar(value):
    if isinstance(value, float) and not math.isfinite(value):
        return {'nonfinite_float': repr(value)}
    return value


class CompleteReader(Reader):
    """Retain stored formula expressions/tokens without executing or translating."""
    def __init__(self, path):
        super().__init__(path)
        self.tokens = {}; self.xml_formulas = {}
        if not self.xls:
            from src.adapters.pdm_workbooks import N
            with zipfile.ZipFile(path) as archive:
                workbook=ET.fromstring(archive.read('xl/workbook.xml'));relationships=ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'));targets={r.attrib['Id']:r.attrib['Target'] for r in relationships}
                for index,sheet in enumerate(workbook.find('s:sheets',N)):
                    target=targets[sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']];target=target.lstrip('/') if target.startswith('/') else 'xl/'+target
                    for cell in ET.fromstring(archive.read(target)).findall('.//s:sheetData/s:row/s:c',N):
                        formula=cell.find('s:f',N)
                        if formula is not None:self.xml_formulas[index,cell.attrib['r']]={'expression':formula.text,'attributes':dict(formula.attrib)}
        if self.xls:
            with olefile.OleFileIO(path) as archive:
                data = archive.openstream('Workbook' if archive.exists('Workbook') else 'Book').read()
            offsets = []; at = 0
            while at + 4 <= len(data):
                kind, size = struct.unpack_from('<HH', data, at); payload = data[at+4:at+4+size]
                if kind == 0x85 and len(payload) >= 4: offsets.append(struct.unpack_from('<I', payload)[0])
                at += size + 4
            for index, start in enumerate(offsets):
                at = start
                while at + 4 <= len(data):
                    kind, size = struct.unpack_from('<HH', data, at); payload = data[at+4:at+4+size]
                    if kind == 6 and size >= 22:
                        row, col = struct.unpack_from('<HH', payload)
                        self.tokens[index, get_column_letter(col+1)+str(row+1)] = base64.b64encode(payload[20:]).decode()
                    at += size + 4
                    if kind == 0x0A: break

    def cell(self,index,address):
        result=super().cell(index,address)
        result['raw_value']=scalar(result['raw_value'])
        return result

    def proof(self, index, address):
        cell = self.cell(index, address)
        result = {k: scalar(v) for k, v in cell.items() if v is not None and k != 'cell'}
        result['raw_value'] = scalar(cell['raw_value'])
        if not result.get('formula_present'):
            result.pop('formula_present', None); result.pop('cache_status', None)
        if result.get('number_format') == 'General': result.pop('number_format')
        if cell['formula_present']:
            result['formula_expression'] = None if self.xls else self.raw[index].get(address, {}).get('formula')
            result['formula_tokens_b64'] = self.tokens.get((index, address))
            if not self.xls:result['formula_xml']=self.xml_formulas.get((index,address),{'expression':None,'attributes':{},'origin':'ARRAY_MASTER_COVERAGE_NOT_TRANSLATED'})
            result['formula_semantics'] = 'STORED_NOT_EXECUTED_OR_TRANSLATED'
        return result

    def populated(self, index):
        rows, cols = self.shape(index)
        if rows * cols > 10_000_000: raise ValueError('SOURCE_SHEET_SCAN_BOUND')
        for row in range(1, rows+1):
            for col in range(1, cols+1):
                address = get_column_letter(col)+str(row)
                cell = self.cell(index, address)
                if cell['raw_value'] is not None or cell['formula_present']:
                    yield address, self.proof(index, address)


def write_complete_plan(sources, catalog, destination, *, as_of, date_rules=None):
    """No DB access. No guessed unit, source asset, date encoding or formula result."""
    dest = Path(destination)
    if dest.exists(): raise ValueError('NEW_COMPLETE_PLAN_REQUIRED')
    dest.mkdir(mode=0o700)
    books, profiles = load_catalog(catalog)
    if len({s['source_ref'] for s in sources}) != len(sources): raise ValueError('DUPLICATE_SOURCE')
    streams = {name: (dest/(name+'.jsonl')).open('x') for name in STREAMS}
    for stream in streams.values(): os.chmod(stream.name, 0o600)
    counts = Counter(); coverage = []; emitted_profiles = set()
    def emit(name, value):
        data = canonical(value)
        if len(data.encode()) > 128*1024: raise ValueError('SOURCE_ITEM_TOO_LARGE')
        streams[name].write(data+'\n'); counts[name] += 1
    try:
        for source in sources:
            book = next(b for b in books if b['id'] == source['source_ref'])
            reader = CompleteReader(source['path']); expected_hash = reader.sha
            source_doc = {'source_ref': book['id'], 'filename': reader.path.name,
                          'sha256': reader.sha, 'size': reader.path.stat().st_size,
                          'sheets': [{'index': i, 'label': name, 'shape': reader.shape(i)} for i, name in enumerate(reader.names)]}
            emit('sources', source_doc)
            try:
                for index, sheet in enumerate(reader.names):
                    physical = dict(reader.populated(index)); claims = defaultdict(set)
                    assigned = set(); local = Counter(); duplicate_claims = Counter()
                    selected = [p for p in profiles if p['source_ref'] == book['id'] and p['sheet_index'] == index]
                    for profile in selected:
                        profile_hash = digest(profile)
                        if profile_hash not in emitted_profiles:
                            emit('profiles', {'profile_hash': profile_hash, 'document': profile}); emitted_profiles.add(profile_hash)
                        for original in iter_events(reader, book, profile, as_of, date_rules=date_rules):
                            refs = []
                            for proof in original['proof']:
                                refs.append({'field_id': proof['field_id'], 'cell': proof['source']['cell'], 'role': 'INPUT'})
                                for context in proof['headers'] + proof['unit_evidence']:
                                    claims[context['cell']].add('CONTEXT')
                            for excluded in original['excluded']:
                                refs.append({'field_id': excluded['field'], 'cell': excluded['cell']['cell'], 'role': excluded['role']})
                            for clock in original['clock_proof']: claims[clock['cell']].add('CONTEXT')
                            present = [r for r in refs if r['cell'] in physical]
                            for ref in refs:
                                if ref['cell'] in physical:
                                    claims[ref['cell']].add(ref['role']); duplicate_claims[ref['cell']] += 1
                            if not present:
                                local['empty_profile_positions_excluded'] += 1
                                continue
                            assigned.update(r['cell'] for r in refs if r['cell'] in physical)
                            reasons = set(original['reasons'])
                            # The staging status is an interpretation gate, never a data deletion gate.
                            reasons.discard('EMPTY_SOURCE_EVENT')
                            if any(isinstance(physical[r['cell']]['raw_value'],dict) and 'nonfinite_float' in physical[r['cell']]['raw_value'] for r in present):reasons.add('NONFINITE_SOURCE_VALUE_RETAINED_RAW')
                            if len({r['cell'] for r in refs}) < len(refs): reasons.add('OVERLAPPING_FIELD_CLAIMS')
                            if any(physical[r['cell']].get('formula_present') for r in present): reasons.add('FORMULA_STORED_NOT_EXECUTED')
                            kind = 'MEASUREMENT_SOURCE' if any(r['role'] == 'INPUT' for r in present) else 'SOURCE_CONTEXT'
                            doc = {'kind': kind, 'classification': 'SOURCE_READY' if original['status'] == 'ACCEPTED' and not reasons else 'RAW_UNRESOLVED',
                                   'reasons': sorted(reasons), 'fields': refs,
                                   'clocks': original['draft']['input_capture']['clocks'],
                                   'clock_cells': [c['cell'] for c in original['clock_proof']],
                                   'source_asset_label': original['draft']['input_capture']['source_asset_label'],
                                   'mapping': 'UNVERIFIED_SOURCE_LABEL_ONLY'}
                            event = {'event_key': original['event_key'], 'source_sha256': reader.sha, 'source_ref': book['id'],
                                     'sheet_index': index, 'profile_id': profile['id'], 'profile_hash': profile_hash,
                                     'axis': original['source']['axis'], 'position': original['source']['position'], 'document': doc}
                            event['classification_hash'] = digest({'parser': PARSER, 'profile_hash': profile_hash,
                                                                  'date_rule': (date_rules or {}).get(profile['id']), 'document': doc})
                            emit('events', event); local[kind] += 1; local[doc['classification']] += 1
                    # All remaining cells are visible source content, not guessed measurements.
                    remaining = defaultdict(list)
                    for address in physical.keys() - assigned:
                        remaining[coordinate_to_tuple(address)[0]].append(address)
                    raw_profile = {'id': 'RAW-ROW-V1', 'version': PARSER, 'source_ref': book['id'], 'sheet_index': index,
                                   'fields': [], 'clocks': [], 'method': 'OTHER', 'approval': 'UNRESOLVED_SOURCE_CONTENT'}
                    raw_hash = digest(raw_profile)
                    if remaining:
                        emit('profiles', {'profile_hash': raw_hash, 'document': raw_profile})
                    for row, addresses in sorted(remaining.items()):
                        addresses.sort(key=lambda a: coordinate_to_tuple(a)[1])
                        refs = [{'field_id': a, 'cell': a, 'role': 'CONTEXT' if claims[a] else 'UNCLASSIFIED'} for a in addresses]
                        doc = {'kind': 'RAW_SOURCE_ROW', 'classification': 'RAW_UNRESOLVED',
                               'reasons': ['UNCLASSIFIED_OR_CONTEXT_SOURCE_CONTENT', 'UNIT_DATE_ASSET_NOT_INFERRED'],
                               'fields': refs, 'clocks': [], 'clock_cells': [], 'source_asset_label': None,
                               'mapping': 'UNVERIFIED_SOURCE_LABEL_ONLY'}
                        event = {'event_key': event_key(reader.sha, index, raw_profile['id'], 'ROW', row),
                                 'source_sha256': reader.sha, 'source_ref': book['id'], 'sheet_index': index,
                                 'profile_id': raw_profile['id'], 'profile_hash': raw_hash, 'axis': 'ROW', 'position': row, 'document': doc}
                        event['classification_hash'] = digest({'parser': PARSER, 'document': doc})
                        emit('events', event); local['RAW_SOURCE_ROW'] += 1
                    category = Counter(); overlaps = 0
                    for address, proof in sorted(physical.items(), key=lambda x: coordinate_to_tuple(x[0])):
                        roles = claims[address]
                        terminal = next((r for r in ('INPUT', 'REFERENCE', 'SOURCE_ASSESSMENT', 'CONTEXT') if r in roles), 'UNCLASSIFIED')
                        category[terminal] += 1
                        if duplicate_claims[address] > 1: overlaps += 1
                        emit('cells', {'source_sha256': reader.sha, 'sheet_index': index, 'cell': address,
                                       'terminal_role': terminal, 'claims': sorted(roles), 'document': proof})
                    if sum(category.values()) != len(physical): raise ValueError('SOURCE_CELL_COVERAGE_MISMATCH')
                    coverage.append({'source_ref': book['id'], 'source_sha256': reader.sha, 'sheet_index': index,
                                     'sheet_label': sheet, 'profiles': len(selected), 'populated_cells': len(physical),
                                     'terminal_cells': dict(category), 'overlapping_value_claim_cells': overlaps,
                                     'events': dict(local), 'catalog_hash_match': reader.sha == book['sha256']})
            finally: reader.close()
            if sha256(Path(source['path']).read_bytes()).hexdigest() != expected_hash: raise ValueError('SOURCE_CHANGED_DURING_SCAN')
        for stream in streams.values(): stream.flush(); os.fsync(stream.fileno()); stream.close()
        implementation=digest([sha256(Path(__file__).read_bytes()).hexdigest(),sha256(Path(__import__('src.adapters.pdm_workbooks',fromlist=['Reader']).__file__).read_bytes()).hexdigest()])
        manifest = {'parser_implementation_sha256':implementation,'schema': 'NADI_COMPLETE_SOURCE_PLAN_V1', 'parser': PARSER, 'as_of': as_of.isoformat(),
                    'catalog_sha256': sha256(Path(catalog).read_bytes()).hexdigest(), 'counts': dict(counts), 'coverage': coverage,
                    'stream_sha256': {n: sha256((dest/(n+'.jsonl')).read_bytes()).hexdigest() for n in STREAMS},
                    'scope': 'PROVISIONAL_SOURCE_NO_CANONICAL_MAPPING', 'database_writes': 0}
        manifest['plan_id'] = digest(manifest)
        p = dest/'manifest.json'; p.write_text(canonical(manifest)+'\n'); p.chmod(0o600)
        return manifest
    except BaseException:
        for stream in streams.values():
            if not stream.closed: stream.close()
        (dest/'INCOMPLETE').write_text('Not importable. Preserve for diagnosis.\n')
        raise


def verify_plan(directory):
    directory = Path(directory)
    if directory.resolve() != directory or directory.is_symlink() or directory.stat().st_mode & 0o077:
        raise ValueError('PRIVATE_COMPLETE_PLAN_REQUIRED')
    def checked(path):
        st = path.lstat()
        if path.is_symlink() or not path.is_file() or st.st_uid != os.getuid() or st.st_mode & 0o077 or st.st_size > 128*1024*1024:
            raise ValueError('PRIVATE_BOUNDED_PLAN_FILE_REQUIRED')
    checked(directory/'manifest.json')
    manifest = json.loads((directory/'manifest.json').read_text()); identity = manifest.pop('plan_id')
    if manifest.get('schema') != 'NADI_COMPLETE_SOURCE_PLAN_V1' or manifest.get('parser') != PARSER or (directory/'INCOMPLETE').exists() or digest(manifest) != identity:
        raise ValueError('COMPLETE_PLAN_IDENTITY_INVALID')
    implementation=digest([sha256(Path(__file__).read_bytes()).hexdigest(),sha256(Path(__import__('src.adapters.pdm_workbooks',fromlist=['Reader']).__file__).read_bytes()).hexdigest()])
    if manifest.get('parser_implementation_sha256')!=implementation:raise ValueError('COMPLETE_PLAN_PARSER_CHANGED')
    for name in STREAMS:
        path = directory/(name+'.jsonl'); checked(path)
        h = sha256()
        with path.open('rb') as stream:
            for block in iter(lambda: stream.read(1024*1024), b''): h.update(block)
        if h.hexdigest() != manifest['stream_sha256'][name]: raise ValueError('COMPLETE_PLAN_CONTENT_CHANGED')
    return dict(manifest, plan_id=identity)
