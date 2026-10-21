"""Private source staging, separate from application/Mart storage and approval."""
from collections import Counter
from datetime import date,datetime
from hashlib import sha256
from pathlib import Path
import json,math,os,re
from openpyxl.utils.cell import coordinate_to_tuple,range_boundaries,get_column_letter
from src.adapters.pdm_workbooks import Reader
from src.domain.manual_inspection import InspectionDraft

def canonical(value):return json.dumps(value,sort_keys=True,separators=(',',':'),ensure_ascii=False)
def fingerprint(value):return sha256(canonical(value).encode()).hexdigest()
def event_key(source_hash,sheet,profile,axis,position):
    # Independent of selected field subset, batch label, and row values.
    return fingerprint([source_hash,sheet,profile,axis,position])

def capture_event_key(capture):
    coords=[coordinate_to_tuple(r['value_region']) for r in capture['rows'] if r.get('value_region')]
    if not coords or any(':' in r.get('value_region','') for r in capture['rows']):raise ValueError('EXISTING_EVENT_COORDINATES_AMBIGUOUS')
    if len({r for r,c in coords})==1:axis,pos='ROW',coords[0][0]
    elif len({c for r,c in coords})==1:axis,pos='COL',coords[0][1]
    else:raise ValueError('EXISTING_EVENT_COORDINATES_AMBIGUOUS')
    # Single-cell partials need an explicit original axis; never guess.
    if len(coords)==1:
        axis=capture.get('source_event_axis')
        if axis not in {'ROW','COL'}:raise ValueError('EXISTING_SINGLE_CELL_AXIS_REQUIRED')
        pos=coords[0][0 if axis=='ROW' else 1]
    return event_key(capture['source_sha256'],capture['sheet_index'],capture['profile_id'],axis,pos)

def load_catalog(path):
    text=Path(path).read_text()
    def read(name):return json.JSONDecoder().raw_decode(text[text.index('=[',text.index('export const '+name))+1:])[0]
    return read('captureSources'),read('captureProfiles')

def unit_evidence_matches(value,unit):
    # Exact, case-sensitive unit expressions only. Exponents/prefixes/suffixes
    # are semantic; do not strip them or infer aliases from substrings.
    norm=lambda v:" ".join(str(v).split())
    text=norm(value);expected=norm(unit)
    candidates=[text,*re.findall(r"[\(\[]([^()\[\]]+)[\)\]]",text)]
    return bool(expected) and any(norm(v)==expected for v in candidates)

def source_clock(cell,label,encoding=None,kind="DATE"):
    if kind=="TIME":
        return {"role":"UNSPECIFIED","raw_value":cell["raw_value"],"precision":"RAW" if cell["raw_value"] is not None else "UNKNOWN","value":None,"source_label":label+" · "+cell["cell"]}
    raw=cell['raw_value'];value=None
    if not cell['formula_present']:
        if cell['date_interpretation']:
            try:
                dt=datetime.fromisoformat(cell['date_interpretation'])
                if dt.year>=1970 and dt.hour==dt.minute==dt.second==0:value=dt.date().isoformat()
            except ValueError:pass
        elif isinstance(raw,str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}',raw):
            try:value=date.fromisoformat(raw).isoformat()
            except ValueError:pass
        elif encoding=='YYYYMMDD' and isinstance(raw,(str,int,float)) and not isinstance(raw,bool):
            text=str(int(raw)) if isinstance(raw,(int,float)) and raw==int(raw) else str(raw)
            try:
                if re.fullmatch(r'\d{8}',text):value=datetime.strptime(text,'%Y%m%d').date().isoformat()
            except ValueError:pass
    return {'role':'UNSPECIFIED','raw_value':raw,'precision':'DATE' if value else 'RAW' if raw is not None else 'UNKNOWN','value':value,'source_label':label+' · '+cell['cell']}

def iter_events(reader,book,profile,as_of,*,date_rules=None,reviewed_hashes=(),existing=()):
    i=profile['sheet_index'];nr,nc=reader.shape(i);fields=profile['fields'];positions=set()
    for f in fields:
        if not f.get('value_region'):continue
        a,b,c,d=range_boundaries(f['value_region'])
        if a==c:positions.update(('ROW',r) for r in range(b,min(d,nr)+1))
        elif b==d:positions.update(('COL',col) for col in range(a,min(c,nc)+1))
    existing_by_key={e['event_key']:e for e in existing}
    if len(existing_by_key)!=len(existing):raise ValueError('DUPLICATE_EXISTING_EVENT_INDEX')
    for axis,pos in sorted(positions):
        reasons=[];rows=[];proof=[];excluded=[];nonempty=0
        if reader.sha!=book['sha256'] and reader.sha not in reviewed_hashes:reasons.append('SOURCE_HASH_CHANGED_REVIEW_REQUIRED')
        for f in fields:
            if not f.get('value_region'):continue
            a,b,c,d=range_boundaries(f['value_region'])
            if axis=='ROW' and (a!=c or not b<=pos<=d):continue
            if axis=='COL' and (b!=d or not a<=pos<=c):continue
            address=get_column_letter(a)+str(pos) if axis=='ROW' else get_column_letter(pos)+str(b)
            cell=reader.cell(i,address)
            if f['row_role']!='INPUT':
                excluded.append({'field':f['id'],'role':f['row_role'],'cell':cell,'reason':'NOT_MEASUREMENT_INPUT'});continue
            raw=cell['raw_value'];nonempty+=raw is not None
            def context_cell(address):
                source=reader.cell(i,address)
                return {'cell':address,'value':source['raw_value'],'formula_present':source['formula_present'],'source_type':source['source_type']}
            headers=[context_cell(h) for h in f['header_cells']]
            unit_proof=[context_cell(u) for u in f.get('unit_cells',[])]
            norm=lambda v:' '.join(str(v).split())
            if cell['formula_present'] or any(x['formula_present'] for x in headers+unit_proof):reasons.append('FORMULA_CACHE_REVIEW_REQUIRED')
            if raw is not None:
                if len(headers)!=1 or norm(headers[0]['value'])!=norm(f['source_label']):reasons.append('HEADER_MISMATCH')
                if not f['unit'] or not f['unit'].strip():reasons.append('UNIT_UNKNOWN')
                elif not any(unit_evidence_matches(u['value'],f['unit']) for u in unit_proof):reasons.append('UNIT_HEADER_MISMATCH')
                if cell['source_type'] in ['e',5]:reasons.append('EXCEL_ERROR')
                if isinstance(raw,float) and not math.isfinite(raw):reasons.append('NONFINITE_SOURCE_VALUE')
            proof.append({'field_id':f['id'],'source':cell,'headers':headers,'unit_evidence':unit_proof})
            rows.append({'row_id':'cell-'+address,'source_field_id':f['id'],'source_label':f['source_label'],'value_region':address,'raw_value':raw,'unit':f['unit'],'header_cells':f['header_cells'],'source_point':f.get('source_point'),'source_axis':f.get('source_axis'),'value_origin':'SOURCE_FORMULA_RESULT' if cell['formula_present'] else 'SOURCE_LITERAL','row_role':'INPUT','contexts':[{'role':'REPORT','label':'Formula/cache status','raw_value':cell['cache_status'],'unit':None}],'clock':None})
        key=event_key(reader.sha,i,profile['id'],axis,pos)
        clock_proof=[];clocks=[]
        rule=(date_rules or {}).get(profile['id'])
        if rule:
            if rule.get('source_sha256')!=reader.sha or rule.get('sheet_index')!=i or rule.get('axis')!=axis or not rule.get('rationale'):reasons.append('DATE_RULE_IDENTITY_MISMATCH');rule=None
        addresses=[]
        for clock in profile['clocks']:
            for header in clock['header_cells']:
                r,c=coordinate_to_tuple(header);address=get_column_letter(c)+str(pos) if axis=='ROW' else get_column_letter(pos)+str(r)
                addresses.append((address,clock['source_label'],clock.get('role','DATE')))
        reviewed_address=None
        if rule:
            locator=str(rule['locator']);reviewed_address=locator+str(pos) if axis=='ROW' else get_column_letter(pos)+locator
            if not any(a==reviewed_address and kind=='DATE' for a,_,kind in addresses):addresses.append((reviewed_address,'Explicit reviewed source date rule','DATE'))
        date_values=[]
        for address,label,kind in dict.fromkeys(addresses):
            cell=reader.cell(i,address)
            encoding=rule.get('encoding') if rule and address==reviewed_address and kind=='DATE' else None
            clock_proof.append(dict(cell,catalog_clock_role=kind,reviewed_date_encoding=encoding))
            parsed=source_clock(cell,label,encoding,kind);clocks.append(parsed)
            if cell['formula_present']:reasons.append('FORMULA_CACHE_REVIEW_REQUIRED')
            if cell['source_type'] in ['e',5]:reasons.append('SOURCE_CLOCK_ERROR_REVIEW_REQUIRED')
            if kind!='TIME':
                if parsed['value'] is not None:date_values.append(parsed['value'])
                elif cell['raw_value'] is not None:reasons.append('SOURCE_DATE_COMPONENT_UNRESOLVED')
        dates=set(date_values)
        if len(dates)!=1:reasons.append('SOURCE_DATE_UNRESOLVED' if not dates else 'SOURCE_DATES_CONFLICT')
        elif next(iter(dates))>as_of.isoformat():reasons.append('FUTURE_SOURCE_DATE_REVIEW_REQUIRED')
        if len(rows)>100:reasons.append('CAPTURE_FIELD_LIMIT')
        if not nonempty:status='REJECTED';reasons.append('EMPTY_SOURCE_EVENT')
        elif any(r in reasons for r in ['EXCEL_ERROR','NONFINITE_SOURCE_VALUE']):status='REJECTED'
        else:status='UNRESOLVED' if reasons else 'ACCEPTED'
        prior=existing_by_key.get(key)
        if prior:
            source={r['value_region']:(r['raw_value'],r['unit']) for r in rows}
            old={r['value_region']:(r['raw_value'],r['unit']) for r in prior['rows']}
            if any(source.get(c)!=v for c,v in old.items()):status='UNRESOLVED';reasons.append('EXISTING_SAMPLE_CONFLICT')
            else:status='EXISTING_PARTIAL' if set(old)!=set(source) else 'EXISTING_COMPLETE';reasons.append('EXISTING_EVENT_PRESERVED_NO_OVERWRITE')
        capture={'schema_version':'1.0','purpose':'PROVISIONAL_INPUT_ONLY','source_ref':book['id'],'source_sha256':reader.sha,'sheet_index':i,'source_sheet_label':reader.names[i],'source_asset_label':reader.names[i] if book['id']!='WB-6KV' else None,'profile_id':profile['id'],'profile_version':profile['version'],'rows':rows,'clocks':clocks}
        if book['id']=='WB-6KV':reasons.append('SOURCE_ASSET_BLOCK_LABEL_REVIEW_REQUIRED');status='UNRESOLVED' if not prior and nonempty else status
        note='DEMO/HISTORIS · event '+key+' · '+reader.path.name+' · sheet '+str(i+1)+' · '+reader.names[i]+'. Source labels unverified; QA asset is an access scaffold, not an operational mapping. Historical performer unknown; author is the demo entry custodian.'
        draft={'canonical_asset_id':'NADI-ASSET-001','method':profile['method'],'method_version':profile['version'],'inspector_ref':'RESOLVE_TRUSTED_AUTHOR','inspected_at':None,'measurements':[],'observations':[note],'operating_context':note,'interpretations':[],'attachments':[],'case_refs':[],'method_extension':{},'input_capture':capture}
        if status=='ACCEPTED':
            try:InspectionDraft.model_validate(draft)
            except ValueError:status='REJECTED';reasons.append('SOURCE_CONTRACT_BOUNDS_INVALID')
        yield {'event_key':key,'request_id':'history:'+key[:56],'status':status,'reasons':sorted(set(reasons)),'source':{'file':reader.path.name,'sha256':reader.sha,'sheet_index':i,'sheet':reader.names[i],'profile':profile['id'],'axis':axis,'position':pos},'proof':proof,'clock_proof':clock_proof,'excluded':excluded,'draft':draft,'existing_record_id':prior.get('record_id') if prior else None}

def write_plan(sources,catalog_path,destination,*,as_of,date_rules=None,reviewed_hashes=(),existing=()):
    dest=Path(destination)
    if dest.exists():raise ValueError('NEW_PRIVATE_PLAN_DIRECTORY_REQUIRED')
    dest.mkdir(mode=0o700,parents=False);books,profiles=load_catalog(catalog_path)
    counts=Counter();reasons=Counter();coverage=[];stream=dest/'events.jsonl'
    try:
        with stream.open('x') as output:
            stream.chmod(0o600)
            for source in sources:
                book=next(b for b in books if b['id']==source['source_ref']);reader=Reader(Path(source['path']));local=Counter();values=Counter()
                try:
                    selected=[p for p in profiles if p['source_ref']==book['id'] and p['sheet_index']<len(reader.names)]
                    for profile in selected:
                        for event in iter_events(reader,book,profile,as_of,date_rules=date_rules,reviewed_hashes=reviewed_hashes,existing=existing):
                            output.write(canonical(event)+'\n');counts[event['status']]+=1;local[event['status']]+=1
                            reasons.update(event['reasons']);values[event['status']]+=sum(r['raw_value'] is not None for r in event['draft']['input_capture']['rows'])
                    coverage.append({'file':reader.path.name,'sha256':reader.sha,'catalog_hash_match':reader.sha==book['sha256'],'sheets':len(reader.names),'covered_sheets':len({p['sheet_index'] for p in selected}),'profiles':len(selected),'events':dict(local),'nonempty_input_values_by_event_status':dict(values)})
                finally:reader.close()
                if sha256(Path(source['path']).read_bytes()).hexdigest()!=reader.sha:raise ValueError('SOURCE_CHANGED_DURING_PLAN')
        manifest={'schema':'NADI_HISTORY_STAGE_V1','as_of':as_of.isoformat(),'sources':coverage,'counts':dict(counts),'reason_counts':dict(reasons),'events_sha256':sha256(stream.read_bytes()).hexdigest(),'catalog_sha256':sha256(Path(catalog_path).read_bytes()).hexdigest(),'date_rules':date_rules or {},'reviewed_hashes':sorted(reviewed_hashes),'existing_event_keys':sorted(e['event_key'] for e in existing),'scope':'PROVISIONAL_DEMO_ONLY_NO_CANONICAL_MAPPING','database_writes':0}
        manifest['plan_id']=fingerprint(manifest);p=dest/'manifest.json';p.write_text(canonical(manifest)+'\n');p.chmod(0o600)
        return manifest
    except Exception:
        # Preserve partial output for diagnosis; never mistake it for a complete plan.
        (dest/'INCOMPLETE').write_text('Plan failed; no manifest is authorized for import.\n');raise

def verify_manifest(plan_dir):
    dest=Path(plan_dir);m=json.loads((dest/'manifest.json').read_text());identity=m.pop('plan_id')
    if fingerprint(m)!=identity or (dest/'INCOMPLETE').exists():raise ValueError('PLAN_IDENTITY_INVALID')
    stream=dest/'events.jsonl'
    digest=sha256()
    with stream.open('rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''):digest.update(block)
    if digest.hexdigest()!=m['events_sha256']:raise ValueError('PLAN_CONTENT_CHANGED')
    m['plan_id']=identity
    return m

def verified_events(plan_dir):
    dest=Path(plan_dir);verify_manifest(dest);stream=dest/'events.jsonl'
    with stream.open() as f:
        for line in f:yield json.loads(line)
